from docx import Document
import openpyxl



class CLGenerator:

	def __init__(self,valSheet,templateDoc,namePrefix,keyval,nameSuffix):
		self.valSheet = valSheet
		self.templateDoc = templateDoc
		self.namePrefix = namePrefix
		self.keyval = keyval
		self.nameSuffix = nameSuffix
		self.replacementKey = []
		self.replacementArray = []

	def replaceAll(self,oldTxt,newTxt):    		
		for p in self.templateDoc.paragraphs:
        		if oldTxt in p.text:
            			inline = p.runs
            			# Loop added to work with runs (strings with same style)
            			for i in range(len(inline)):
            			    if oldTxt in inline[i].text:
            			        text = inline[i].text.replace(oldTxt, newTxt)
            			        inline[i].text = text
    		return 1

	def replaceSaveRevert(self,oldTxtLst,newTxtLst):
		for i in range(len(oldTxtLst)): #replace all values
			self.replaceAll(oldTxtLst[i],newTxtLst[i])
		self.templateDoc.save(self.namingConvention(newTxtLst)) #save with unique name
		print(self.namingConvention(newTxtLst) + " finished")
		for i in range(len(oldTxtLst)): #revert
			self.replaceAll(newTxtLst[i],oldTxtLst[i])
			
	def namingConvention(self,lst):
		return self.namePrefix + lst[self.keyval] + self.nameSuffix + '.docx'

	def generateBatch(self):
		tmpLst = []
		for i in range(len(self.replacementArray[0])):
			for j in range(len(self.replacementKey)):
				tmpLst.append(self.replacementArray[j][i])
			self.replaceSaveRevert(self.replacementKey,tmpLst)
			while len(tmpLst) > 0:
				tmpLst.pop()
		

	#oldtext = str (that will be replaced from template file
	#columnstart = char (letter of column replacement list is in)
	#rowstart = int row number that list starts on
	def addReplacementList(self,oldtext,columnstart,rowstart):
		retVal = []
		currentCell = columnstart + str(rowstart)
		tmpInt = 0
		while (self.valSheet[currentCell].value != None):
			retVal.append(self.valSheet[currentCell].value)
			tmpInt = tmpInt + 1
			currentCell = columnstart + str(rowstart+tmpInt)
		self.replacementArray.append(retVal)
		self.replacementKey.append(oldtext)



wb = openpyxl.load_workbook('joblist.xlsx')
sheet = wb.get_sheet_by_name('Sheet1')
template = Document('coverlettertemplate.docx')
replacer = CLGenerator(sheet,template,'FirstnameMcLastname',0,'CoverLetter')
replacer.addReplacementList('THATCOMPANY','A',3)
replacer.addReplacementList('THATPOSITION','B',3)
replacer.addReplacementList('SKILL1','C',3)
replacer.addReplacementList('SKILL2','D',3)
replacer.generateBatch()

